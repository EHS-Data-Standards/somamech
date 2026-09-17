"""Tests for the pooled-workbook generator (scripts/build_workbook.py)."""
import shutil
import subprocess
from pathlib import Path

import openpyxl

ROOT = Path(__file__).parent.parent
SCRIPT = ROOT / "scripts" / "build_workbook.py"
VALID = ROOT / "tests" / "data" / "valid"


def run(*args):
    return subprocess.run(
        ["uv", "run", "python", str(SCRIPT), *args],
        capture_output=True, text=True, cwd=ROOT,
    )


def test_pools_multiple_papers_into_one_workbook(tmp_path):
    kb = tmp_path / "kb"
    kb.mkdir()
    shutil.copy(VALID / "Container-liu2024-pm25-cftr.yaml", kb)
    shutil.copy(VALID / "Container-montgomery2020-pm25-mucociliary.yaml", kb)
    out = tmp_path / "pooled.xlsx"

    result = run("--kb-dir", str(kb), "--output", str(out))
    assert result.returncode == 0, result.stdout + result.stderr
    assert out.exists()

    wb = openpyxl.load_workbook(out)
    assert "Papers" in wb.sheetnames
    assert wb["Papers"].max_row - 1 == 2  # one row per paper
    # every data tab's first column names the source paper
    for ws in wb.worksheets:
        assert ws.cell(1, 1).value == "source_publication"
    # rows from both papers land in the same GeneExpressionAssay tab
    sources = {r[0] for r in wb["GeneExpressionAssay"].iter_rows(min_row=2, values_only=True)}
    assert len(sources) == 2


def test_cross_paper_id_collision_fails(tmp_path):
    kb = tmp_path / "kb"
    kb.mkdir()
    original = (VALID / "Container-liu2024-pm25-cftr.yaml").read_text()
    (kb / "a.yaml").write_text(original)
    # same content, different declared paper -> every shared assay id collides
    (kb / "b.yaml").write_text(original.replace("PMID:39113893", "PMID:99999999"))

    result = run("--kb-dir", str(kb), "--check-only")
    assert result.returncode != 0
    assert "must be unique to one paper" in result.stderr


def test_empty_kb_dir_is_fine(tmp_path):
    result = run("--kb-dir", str(tmp_path), "--check-only")
    assert result.returncode == 0, result.stdout + result.stderr
