"""Tests for yaml_to_excel.py converter."""
import datetime
import os
import subprocess
import sys
import tempfile

import pytest

SCRIPT = os.path.join(os.path.dirname(__file__), "..", "scripts", "yaml_to_excel.py")
DATA_DIR = os.path.join(os.path.dirname(__file__), "data", "valid")

EXAMPLE_YAMLS = [
    (
        "Container-liu2024-pm25-cftr.yaml",
        {
            "required_tabs": [
                "Metadata", "Protocol", "ExposureCondition", "KeyEvent",
                "CellularSystem", "InVivoSubject",
                "CFTRFunctionAssay", "CFTRFunctionOutput",
                "GeneExpressionAssay", "GeneExpressionOutput",
                "GobletCellAssay", "GobletCellOutput",
                "BALFSputumAssay", "BALFSputumOutput",
                "LungFunctionAssay", "LungFunctionOutput",
                "Responses", "ResponseComparison", "KeyEventRelationship",
            ],
            "min_assay_rows": {
                # Calu-3 control/PM2.5 Isc pair merged into one assay with
                # one output record per condition
                "CFTRFunctionAssay": 3,
                "CFTRFunctionOutput": 4,
                "GeneExpressionAssay": 5,
                "GobletCellAssay": 2,
                "LungFunctionAssay": 2,
                "ResponseComparison": 2,
            },
        },
    ),
    (
        "Container-montgomery2020-pm25-mucociliary.yaml",
        {
            "required_tabs": [
                "Metadata", "Protocol", "ExposureCondition", "KeyEvent",
                "CellularSystem",
                "GeneExpressionAssay", "GeneExpressionOutput",
                "GobletCellAssay", "GobletCellOutput",
                "FoxJExpressionAssay", "FoxJExpressionOutput",
            ],
            "min_assay_rows": {
                "GeneExpressionAssay": 6,
                "GobletCellAssay": 4,
                "FoxJExpressionAssay": 2,
            },
        },
    ),
    (
        "Container-response_comparison.yaml",
        {
            "required_tabs": [
                "ExposureCondition", "KeyEvent", "InVivoSubject",
                "LungFunctionAssay", "LungFunctionOutput",
                "Responses", "ResponseComparison", "KeyEventRelationship",
            ],
            "min_assay_rows": {
                # one output record per experimental condition
                "LungFunctionOutput": 2,
                # long format: one row per condition x measurement
                "Responses": 2,
                "ResponseComparison": 1,
            },
        },
    ),
]


@pytest.mark.parametrize("yaml_name,expectations", EXAMPLE_YAMLS, ids=[e[0] for e in EXAMPLE_YAMLS])
def test_yaml_to_excel(yaml_name, expectations):
    """Test that yaml_to_excel.py produces a valid Excel file with expected tabs and rows."""
    import openpyxl

    yaml_path = os.path.join(DATA_DIR, yaml_name)
    assert os.path.exists(yaml_path), f"YAML file not found: {yaml_path}"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        output_path = tmp.name

    try:
        result = subprocess.run(
            ["uv", "run", "python", SCRIPT, "--input", yaml_path, "--output", output_path],
            capture_output=True, text=True, timeout=60,
        )
        assert result.returncode == 0, f"Script failed:\nstdout: {result.stdout}\nstderr: {result.stderr}"
        assert os.path.exists(output_path), "Output file was not created"

        wb = openpyxl.load_workbook(output_path)
        sheet_names = wb.sheetnames

        # Check required tabs exist
        for tab in expectations["required_tabs"]:
            assert tab in sheet_names, f"Missing tab: {tab}. Found: {sheet_names}"

        # Check minimum row counts (header row + data rows)
        for tab, min_rows in expectations.get("min_assay_rows", {}).items():
            ws = wb[tab]
            data_rows = ws.max_row - 1  # subtract header
            assert data_rows >= min_rows, (
                f"Tab '{tab}' has {data_rows} data rows, expected >= {min_rows}"
            )

        # Verify no empty tabs for collections that should have data
        for tab in expectations["required_tabs"]:
            ws = wb[tab]
            assert ws.max_row >= 2, f"Tab '{tab}' appears empty (only header row)"

        wb.close()
    finally:
        if os.path.exists(output_path):
            os.unlink(output_path)


# ---------------------------------------------------------------------------
# _fmt_cell: coercion for header-driven slots whose range is an inlined class
# ---------------------------------------------------------------------------

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))

import yaml_to_excel as y2e  # noqa: E402 - needs the sys.path line above


@pytest.mark.parametrize(
    "value,expected",
    [
        # scalars pass through untouched -- openpyxl writes these itself
        ("plain string", "plain string"),
        (42, 42),
        (1.5, 1.5),
        (True, True),
        (None, None),
        ("", ""),
        # openpyxl writes temporal types as real typed cells, so they pass
        # through rather than being flattened into text
        (datetime.date(2021, 1, 1), datetime.date(2021, 1, 1)),
        (datetime.datetime(2021, 1, 1, 9, 30), datetime.datetime(2021, 1, 1, 9, 30)),
        (datetime.time(9, 30), datetime.time(9, 30)),
        (datetime.timedelta(days=1), datetime.timedelta(days=1)),
        # an ontology term reads the way the hand-written row builders write one
        ({"id": "CL:0000771", "name": "eosinophil"}, "eosinophil (CL:0000771)"),
        ({"id": "CL:0000775"}, "CL:0000775"),
        ({"name": "eosinophil"}, "eosinophil"),
        # a measurement keeps its unit
        ({"value": "10", "unit": {"id": "UO:0000275", "name": "ng/mL"}}, "10 ng/mL (UO:0000275)"),
        # a dict carrying both an id and a measurement keeps the measurement
        ({"id": "X:1", "value": "10", "unit": {"name": "ng/mL"}}, "10 ng/mL"),
        # a unit with no value still reads as a label, not a repr
        ({"unit": {"id": "UO:1", "name": "ng/mL"}}, "ng/mL (UO:1)"),
        # a present-but-null key blanks instead of rendering the text 'None'
        ({"value": None}, ""),
        ({"unit": None}, ""),
        ({"value": None, "unit": {"name": "ng/mL"}}, "ng/mL"),
        # ...but a real zero is a measurement, not an absence
        ({"value": 0, "unit": {"name": "ng/mL"}}, "0 ng/mL"),
        ({"value": "0"}, "0"),
        # anything else flattens rather than reaching openpyxl as an object
        (["a", "b"], "a; b"),
        ([{"gene": "Foxj1"}], "gene: Foxj1"),
        ({"other": "shape"}, "other: shape"),
    ],
)
def test_fmt_cell_coerces_to_something_excel_accepts(value, expected):
    assert y2e._fmt_cell(value) == expected


def test_assay_row_never_yields_a_non_scalar():
    """The regression that broke `just generate-workbook`: target_cell_type is
    an inlined CellTypeReference, so the generic branch handed openpyxl a dict.
    """
    assay = {
        "id": "ASSAY:x",
        "target_cell_type": {"id": "CL:0000771", "name": "eosinophil"},
    }
    row = y2e._assay_row(assay, y2e.HEADERS["BALFSputumAssay"])
    assert all(isinstance(c, (str, int, float, bool, type(None))) for c in row), row
    assert "eosinophil (CL:0000771)" in row


@pytest.mark.parametrize(
    "fixture_name",
    sorted(f for f in os.listdir(DATA_DIR) if f.endswith(".yaml")),
)
def test_every_valid_fixture_converts(fixture_name, tmp_path):
    """Every committed valid fixture must survive the converter.

    `just pipeline-test` runs two of them by name, so a fixture it does not
    name could stop converting without anything failing -- which is how the
    target_cell_type crash reached main.
    """
    out = tmp_path / "out.xlsx"
    result = subprocess.run(
        ["uv", "run", "python", SCRIPT,
         "--input", os.path.join(DATA_DIR, fixture_name),
         "--output", str(out)],
        capture_output=True, text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr
    assert out.exists()
